"""Read-only adapters and deterministic reconciliation for the two EM admin sheets."""
from __future__ import annotations

import hashlib
import io
import json
import re
import unicodedata
import zipfile
from datetime import date, datetime
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.parse import urlencode
from urllib.request import Request, urlopen
import os

from .domain import FIELDS, normalize_method
from .importers import MAX_XLSX_UNCOMPRESSED_BYTES

SPREADSHEET_ID = "1Bh1uNUIVBSQ2IuY2M5o_CDuLfqlPBsH35S_nyHg8XX8"
SHEETS = {
    "三廠管理者更新": {"sheet_id": 0, "site": "三廠", "columns": 25},
    "一廠管理者更新": {"sheet_id": 526085303, "site": "一廠", "columns": 26},
}
HEADER_ROW = [
    "三廠/一廠環測製測 Data", "批次編號/點位", "年", "月", "日", "檢測方式",
    "檢測目的", "檢驗批次", "操作者", "操作者2", "科室", "Grade", "點位",
    "總菌數 (CFU)", "菌落數 (CFU)", "PA/AL/AC classification", "備註", "F",
    "菌鑑結果流水號分群", "菌鑑方式", "菌種名稱", "Sorce", "送鑑單位", "操作人員", "備註",
]
SPECIES_START = 18  # S:Y are species evidence, not additional samples.


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def _row(values, title, row_number, formula_cells=None):
    values = list(values[:25])
    values.extend([None] * (25 - len(values)))
    return {"row_number": row_number, "site": SHEETS[title]["site"],
            "sheet_title": title, "sheet_id": None,
            "values": [_text(value) for value in values],
            "formula_cells": formula_cells or {}}


def parse_xlsx_targets(content: bytes, *, complete: bool = False) -> dict:
    """Read only the two named target tabs; never iterate unrelated worksheet rows."""
    if not content or len(content) > 25 * 1024 * 1024:
        raise ValueError("XLSX 空白或超過 25 MB；未讀取。")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise ValueError("XLSX 解壓縮後超過 80 MB；未讀取。")
            if any(item.filename.lower().endswith("vbaproject.bin") for item in archive.infolist()):
                raise ValueError("不支援含 VBA 巨集的 XLSX。")
        from openpyxl import load_workbook
        formula_workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        value_workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, OSError) as error:
        raise ValueError("無法讀取有效的 XLSX。") from error
    missing = set(SHEETS) - set(formula_workbook.sheetnames)
    if missing:
        formula_workbook.close()
        value_workbook.close()
        raise ValueError("XLSX 缺少指定分頁：" + "、".join(sorted(missing)))
    rows_by_sheet = {}
    tab_evidence = {}
    try:
        for title in SHEETS:
            worksheet = formula_workbook[title]
            cached_worksheet = value_workbook[title]
            rows = []
            last_row = 1
            formula_rows = worksheet.iter_rows(min_col=1, max_col=25)
            cached_rows = cached_worksheet.iter_rows(min_col=1, max_col=25)
            for number, (formula_cells, cached_cells) in enumerate(zip(formula_rows, cached_rows), 1):
                values, formula_evidence = [], {}
                for index, (formula_cell, cached_cell) in enumerate(zip(formula_cells, cached_cells)):
                    if formula_cell.data_type == "f" or (isinstance(formula_cell.value, str) and
                                                          formula_cell.value.startswith("=")):
                        cached_value = cached_cell.value
                        values.append(cached_value)
                        formula_evidence[chr(65 + index)] = {
                            "formula": _text(formula_cell.value),
                            "cached_value": _text(cached_value) if cached_value is not None else None,
                        }
                    else:
                        values.append(formula_cell.value)
                item = _row(values, title, number, formula_evidence)
                if any(item["values"]):
                    rows.append(item)
                    last_row = number
            if not rows or not _is_header(rows[0]["values"]):
                raise ValueError(f"{title} 缺少可辨識的標題列；未建立快照。")
            tab_evidence[title] = {"sheet_id": None, "range": f"A1:Y{last_row}",
                                  "formula_cells": {str(row["row_number"]): row["formula_cells"]
                                                    for row in rows if row["formula_cells"]}}
            rows_by_sheet[title] = rows
    finally:
        formula_workbook.close()
        value_workbook.close()
    for title, rows in rows_by_sheet.items():
        for row in rows:
            row["sheet_id"] = tab_evidence[title]["sheet_id"]
    return {"rows_by_sheet": rows_by_sheet, "complete": bool(complete), "source": "xlsx",
            "evidence": {"source_kind": "xlsx", "spreadsheet_id": "",
                         "completeness_basis": "使用者標記完整" if complete else "未確認完整",
                         "tabs": tab_evidence}}


def parse_api_ranges(ranges: dict[str, list[list[object]]], *, complete: bool, evidence=None) -> dict:
    """Normalize API values for the same exact tabs and same row adapter as XLSX."""
    if set(ranges) != set(SHEETS):
        raise ValueError("Google API 回應必須且只能包含指定的兩個分頁。")
    rows_by_sheet = {}
    for title in SHEETS:
        tab = (evidence or {}).get("tabs", {}).get(title, {})
        rows = [_row(values, title, index) for index, values in enumerate(ranges[title], 1)]
        for row in rows:
            row["sheet_id"] = tab.get("sheet_id")
        rows = [row for row in rows if any(row["values"])]
        if not rows or not _is_header(rows[0]["values"]):
            raise ValueError(f"{title} 缺少可辨識的標題列；未建立快照。")
        rows_by_sheet[title] = rows
    return {"rows_by_sheet": rows_by_sheet, "complete": bool(complete), "source": "google-api",
            "evidence": evidence or {"source_kind": "google-api", "spreadsheet_id": "",
                                      "completeness_basis": "未提供讀取證據", "tabs": {}}}


def _is_header(values):
    accepted = {
        1: {"批次編號/點位"}, 2: {"年", "年份"}, 3: {"月"}, 4: {"日"},
        5: {"檢測方式", "採樣方式"}, 6: {"檢測目的"}, 7: {"檢驗批次"},
        8: {"操作者"}, 9: {"操作者2"}, 10: {"科室"}, 11: {"Grade"},
        12: {"點位"}, 13: {"總菌數 (CFU)"}, 14: {"菌落數 (CFU)"},
        18: {"菌鑑結果流水號分群"}, 19: {"菌鑑方式"}, 20: {"菌種名稱"},
        21: {"Sorce", "Source"}, 22: {"送鑑單位"}, 23: {"操作人員"}, 24: {"備註"},
    }
    return all(index < len(values) and
               unicodedata.normalize("NFKC", _text(values[index])) in
               {unicodedata.normalize("NFKC", alias) for alias in aliases}
               for index, aliases in accepted.items())


def snapshot_digest(snapshot: dict) -> str:
    rows = {title: [{"row_number": row["row_number"], "site": row["site"], "values": row["values"]}
                    for row in sheet_rows]
            for title, sheet_rows in snapshot["rows_by_sheet"].items()}
    content = json.dumps(rows, ensure_ascii=False, sort_keys=True,
                         separators=(",", ":")).encode("utf-8")
    return hashlib.sha256(content).hexdigest()


def _iso_date(row):
    values = row["values"]
    try:
        y, m, d = (int(_text(values[i])) for i in (2, 3, 4))
        return date(y, m, d).isoformat()
    except (ValueError, TypeError):
        return ""


def _integer(value):
    text = unicodedata.normalize("NFKC", _text(value))
    return str(int(text)) if re.fullmatch(r"[0-9]+", text) else None


def _monitoring_type(purpose):
    return {"環測": "例行環測", "汙染後環測": "污染後環測"}.get(_text(purpose), "")


def _row_identity(row):
    values = row["values"]
    return {
        "site": row["site"], "sample_date": _iso_date(row),
        "method": normalize_method(values[5]), "monitoring_type": _monitoring_type(values[6]),
        "batch_id": _text(values[1]), "batch_no": _text(values[7]),
        "operator_1": _text(values[8]), "operator_2": _text(values[9]),
        "room": _text(values[10]), "grade": _text(values[11]), "point_id": _text(values[12]),
    }


def _missing_identity_reasons(identity):
    reasons = []
    if not identity["batch_id"]:
        reasons.append("B 批次編號/點位空白，不能判定採樣身分")
    if not identity["sample_date"]:
        reasons.append("C/D/E 日期缺漏或不合法")
    if not identity["method"]:
        reasons.append("F 方法空白")
    if not identity["point_id"]:
        reasons.append("M 點位空白")
    return reasons


def _could_share_sample_context(left, right):
    """Return true only when same-site/same-B rows have no explicit identity conflict."""
    if (left["site"] != right["site"] or not left["batch_id"] or
            left["batch_id"] != right["batch_id"]):
        return False
    return all(not left[field] or not right[field] or left[field] == right[field]
               for field in ("sample_date", "method", "batch_no", "room", "grade", "point_id"))


def _pending_row_evidence(row, identity):
    values = row["values"]
    species = [{"cell": f"{chr(65 + index)}{row['row_number']}", "raw_value": values[index]}
               for index in range(SPECIES_START, 25) if values[index]]
    return {
        "tab": row.get("sheet_title", ""), "sheet_id": row.get("sheet_id"),
        "row_number": row["row_number"],
        "source_range": f"A{row['row_number']}:Y{row['row_number']}",
        "missing_identity_fields": [name for name, value in (
            ("C/D/E sample date", identity["sample_date"]), ("F method", identity["method"]),
            ("B batch ID", identity["batch_id"]), ("M point ID", identity["point_id"]),
            ("H test batch", identity["batch_no"]), ("K room", identity["room"]),
            ("L Grade", identity["grade"])) if not value],
        "identity": identity,
        "N_total_cfu": values[13], "O_colony_count": values[14],
        "species_cells_S_Y": species, "cells_A_Y": values,
    }


CFU_RULE_VERSION = "em008-cfu-v2"
FIELD_SOURCE_COLUMNS = {
    "batch_id": (1,), "sample_date": (2, 3, 4), "monitoring_type": (6,), "method": (5,),
    "batch_no": (7,), "operator": (8, 9), "room": (10,), "grade": (11,),
    "point_id": (12,), "result_raw": (13,), "result_type": (13,), "cfu_count": (13,),
}


def field_source_trace(group):
    """Map each proposed value to its source tab, row, and column(s)."""
    sample = group.get("sample") or {}
    trace = {}
    for field, indexes in FIELD_SOURCE_COLUMNS.items():
        if field == "cfu_count" and group.get("zero_inference"):
            indexes = (13, 14)
        cells = []
        for row in group["rows"]:
            for index in indexes:
                column = chr(65 + index)
                cells.append({"tab": row.get("sheet_title", ""), "sheet_id": row.get("sheet_id"),
                              "row": row["row_number"], "cell": f"{column}{row['row_number']}",
                              "raw_value": row["values"][index],
                              "formula": row.get("formula_cells", {}).get(column)})
        trace[field] = {"value_used": sample.get(field, group["identity"].get(field, "")), "cells": cells}
    trace["site"] = {"value_used": group["identity"].get("site", ""),
                     "rule": "指定分頁的廠別脈絡",
                     "tabs": [{"tab": title, "sheet_id": sheet_id} for title, sheet_id in sorted(
                         {(row.get("sheet_title", ""), row.get("sheet_id")) for row in group["rows"]})]}
    trace["monitoring_type"]["rule"] = "G 精確值映射" if sample else "待核對原始目的"
    trace["result_type"]["rule"] = "有效非負整數總 CFU → count" if sample else "待核對"
    if group.get("zero_inference"):
        trace["cfu_count"]["rule"] = group["zero_inference"]
        trace["cfu_count"]["rule_version"] = group.get("zero_rule_version", CFU_RULE_VERSION)
    if group.get("related_pending_rows"):
        trace["related_pending_rows"] = group["related_pending_rows"]
    return trace


def _fingerprint(identity):
    key = {name: identity[name] for name in (
        "site", "sample_date", "method", "room", "grade", "point_id", "batch_id", "batch_no")}
    value = json.dumps(key, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def sample_groups(snapshot: dict) -> list[dict]:
    """Group by sample identity first, then assess purpose and all result evidence."""
    groups = {}
    loose_rows = []
    rows_by_site_batch = {}
    species_rows = 0
    for title in SHEETS:
        for row in snapshot["rows_by_sheet"].get(title, []):
            values = row["values"]
            if _is_header(values) or not any(values):
                continue
            identity = _row_identity(row)
            if identity["batch_id"]:
                rows_by_site_batch.setdefault((identity["site"], identity["batch_id"]), []).append(row)
            species = [value for value in values[SPECIES_START:25] if value]
            species_rows += bool(species)
            reasons = _missing_identity_reasons(identity)
            if reasons:
                loose_rows.append({"row": row, "identity": identity, "species": species,
                                   "reasons": reasons})
                continue
            identity["monitoring_type"] = ""
            fingerprint = _fingerprint(identity)
            group = groups.setdefault(fingerprint, {
                "fingerprint": fingerprint, "identity": identity, "rows": [],
                "species_rows": 0, "reasons": [], "sample": None, "status": "pending",
            })
            group["rows"].append(row)
            group["species_rows"] += bool(species)
            purpose = _text(values[6])
            if purpose:
                group.setdefault("purposes", set()).add(purpose)
    output = list(groups.values())
    for item in loose_rows:
        identity, row = item["identity"], item["row"]
        fingerprint = "pending:" + hashlib.sha256(
            f"{row['site']}:{row['row_number']}:{row['values']}".encode("utf-8")).hexdigest()
        output.append({"fingerprint": fingerprint, "identity": identity, "rows": [row],
                       "species_rows": int(bool(item["species"])), "reasons": item["reasons"],
                       "sample": None, "status": "pending"})

    # A reused B across different dates/methods/points is not merged, and both groups stay pending.
    by_batch = {}
    for group in output:
        identity = group["identity"]
        if identity["batch_id"] and not group["fingerprint"].startswith("pending:"):
            context = (identity["site"], identity["sample_date"], identity["method"],
                       identity["point_id"], identity["room"], identity["grade"], identity["batch_no"])
            by_batch.setdefault((identity["site"], identity["batch_id"]), set()).add(context)
    for group in output:
        identity = group["identity"]
        key = (identity["site"], identity["batch_id"])
        if identity["batch_id"] and len(by_batch.get(key, ())) > 1:
            group["reasons"].append("相同 B 出現在不同日期／方法／點位／房間／Grade／檢驗批次；分開保留並待核對")

    for group in output:
        if group["reasons"]:
            group["status"] = "pending"
            continue
        rows = group["rows"]
        values = [row["values"] for row in rows]
        purposes = group.get("purposes", set())
        mapped_purposes = {_monitoring_type(purpose) for purpose in purposes if _monitoring_type(purpose)}
        unknown_purposes = purposes - {"環測", "汙染後環測"}
        if len(mapped_purposes) > 1 or unknown_purposes:
            group["reasons"].append("G 目的在同一採樣群內矛盾或未有明確監測類型對照：" +
                                    "、".join(sorted(purposes)))
            group["status"] = "pending"
            group["metric_evidence"] = {"N": [row[13] for row in values], "O": [row[14] for row in values],
                                        "species_rows": group["species_rows"]}
            continue
        if len(mapped_purposes) == 1:
            group["identity"]["monitoring_type"] = next(iter(mapped_purposes))
        elif not purposes:
            group["reasons"].append("G 目的未有明確監測類型對照")
            group["status"] = "pending"
            group["metric_evidence"] = {"N": [row[13] for row in values], "O": [row[14] for row in values],
                                        "species_rows": group["species_rows"]}
            continue
        missing_cache = [f"{row['sheet_title']}!{column}{row['row_number']}"
                         for row in rows for column in ("N", "O")
                         if column in row.get("formula_cells", {}) and
                         row["formula_cells"][column].get("cached_value") is None]
        if missing_cache:
            group["reasons"].append("N/O 公式缺少快取計算值，無法判定：" + "、".join(missing_cache))
            group["status"] = "pending"
            group["metric_evidence"] = {"N": [row[13] for row in values], "O": [row[14] for row in values],
                                        "species_rows": group["species_rows"]}
            continue
        total_values = { _integer(row[13]) for row in values if _text(row[13]) }
        invalid_total = any(_text(row[13]) and _integer(row[13]) is None for row in values)
        if invalid_total or len(total_values) > 1:
            group["reasons"].append("N 總菌數不是一致的非負整數")
            continue
        total = next(iter(total_values)) if total_values else None
        raw_total = next((_text(row[13]) for row in values if _text(row[13])), "")
        colony_values = { _integer(row[14]) for row in values if _text(row[14]) }
        invalid_colony = any(_text(row[14]) and _integer(row[14]) is None for row in values)
        child_rows = sum(1 for row in values if any(row[SPECIES_START:25]))
        if total == "0" and (invalid_colony or (colony_values and colony_values != {"0"}) or child_rows):
            group["reasons"].append("N 明確為 0，但 O 或菌鑑子列顯示相反證據")
            continue
        if total is None:
            if invalid_colony or len(colony_values) > 1:
                group["reasons"].append("N 空白且 O 菌落數互相矛盾或非數值")
                continue
            if (not snapshot.get("complete") or invalid_colony or colony_values != {"0"} or child_rows):
                group["reasons"].append("N 空白時不能證明總 CFU=0：需完整快照、O=0、無菌種子列且無矛盾")
                continue
            row_keys = {(row.get("sheet_title", ""), row["row_number"]) for row in rows}
            related = []
            for source_row in rows_by_site_batch.get((group["identity"]["site"],
                                                       group["identity"]["batch_id"]), ()):
                key = (source_row.get("sheet_title", ""), source_row["row_number"])
                if key in row_keys:
                    continue
                other_identity = _row_identity(source_row)
                if _could_share_sample_context(group["identity"], other_identity):
                    related.append(_pending_row_evidence(source_row, other_identity))
            if related:
                group["related_pending_rows"] = related
                group["metric_evidence"] = {
                    "N": [row["values"][13] for row in rows] +
                         [item["N_total_cfu"] for item in related],
                    "O": [row["values"][14] for row in rows] +
                         [item["O_colony_count"] for item in related],
                    "species_cells_S_Y": [cell for row in rows
                                           for cell in _pending_row_evidence(row, _row_identity(row))["species_cells_S_Y"]] +
                                          [cell for item in related for cell in item["species_cells_S_Y"]],
                    "related_source_rows": related,
                }
                details = []
                for item in related:
                    species = ", ".join(cell["raw_value"] for cell in item["species_cells_S_Y"])
                    details.append(
                        f"{item['tab']} row {item['row_number']} "
                        f"N={item['N_total_cfu'] or '空白'}、O={item['O_colony_count'] or '空白'}" +
                        (f"、菌鑑={species}" if species else ""))
                group["reasons"].append(
                    "同廠同 B 有欄位不完整且可能相關的來源列；未合併並保留待核，不能推定 CFU=0：" +
                    "；".join(details))
                continue
            total = "0"
            group["zero_inference"] = "N 原空白；完整快照中 O=0，無菌種子列且無矛盾；總 CFU 候選 0"
            group["zero_rule_version"] = CFU_RULE_VERSION
        operators = {value for row in values for value in (row[8], row[9]) if value}
        operator = next(iter(operators)) if len(operators) == 1 else ""
        if len(operators) > 1:
            group["reasons"].append("I/J 有多位操作者；保留原文並待核對人員歸屬")
            continue
        identity = group["identity"]
        mapped = {field: "" for field in FIELDS}
        mapped.update(sample_date=identity["sample_date"], site=identity["site"],
                      monitoring_type=identity["monitoring_type"], method=identity["method"],
                      room=identity["room"], grade=identity["grade"], point_id=identity["point_id"],
                      operator=operator, batch_no=identity["batch_no"], result_raw=raw_total,
                      result_type="count", cfu_count=total)
        group["sample"] = mapped
        group["status"] = "new"
    for group in output:
        group["field_sources"] = field_source_trace(group)
    return sorted(output, key=lambda item: (item["identity"].get("site", ""),
                                             item["identity"].get("sample_date", ""),
                                             item["identity"].get("point_id", ""),
                                             item["identity"].get("batch_id", ""), item["fingerprint"]))


def match_word_candidates(groups: list[dict], records: list[dict], adopted: dict | None = None):
    """Match only unique Word candidates and suggest fills for still-empty fields."""
    adopted = adopted or {}
    adopted_records = {record["id"]: record for record in records}

    def mismatches(sheet, current):
        different = []
        for field, value in sheet.items():
            existing = current.get(field)
            if value and _text(existing):
                left = normalize_method(existing) if field == "method" else _text(existing)
                right = normalize_method(value) if field == "method" else _text(value)
                if left != right:
                    different.append(field)
        return different

    def adopted_batch(record):
        try:
            source = json.loads(record.get("source_text") or "{}")
            value = _text(source.get("sample_identity", {}).get("batch_id"))
            if value:
                return value
            batches = {_text(row.get("cells_A_Y", [""] * 2)[1]) for row in source.get("rows", [])
                       if len(row.get("cells_A_Y", [])) > 1 and _text(row["cells_A_Y"][1])}
            return next(iter(batches)) if len(batches) == 1 else ""
        except (TypeError, ValueError):
            return ""

    repeats = {}
    for group in groups:
        i = group["identity"]
        if all(i.get(k) for k in ("site", "sample_date", "method", "point_id")):
            core = (i["site"], i["sample_date"], i["method"], i["point_id"])
            repeats.setdefault(core, set()).add(i["batch_id"])
    for group in groups:
        if not group.get("sample") or group["reasons"]:
            continue
        if group["fingerprint"] in adopted:
            group["status"] = "already_adopted"
            group["record_id"] = adopted[group["fingerprint"]]
            record = adopted_records.get(group["record_id"])
            if record:
                differences = mismatches(group["sample"], record["current"])
                if differences:
                    group["status"] = "conflict"
                    group["reasons"].append("已採用樣本的新來源值與目前紀錄衝突：" + ", ".join(differences))
            continue
        sheet = group["sample"]
        base_candidates = []
        for record in records:
            current = record["current"]
            if record.get("origin_kind") == "sheet":
                known_batch = adopted_batch(record)
                if known_batch and known_batch != group["identity"].get("batch_id"):
                    continue
            if (current.get("sample_date") != sheet["sample_date"] or
                    normalize_method(current.get("method")) != sheet["method"] or
                    _text(current.get("point_id")) != sheet["point_id"]):
                continue
            base_candidates.append(record)
        core = (sheet["site"], sheet["sample_date"], sheet["method"], sheet["point_id"])
        word_candidates = [record for record in base_candidates
                           if record.get("origin_kind") != "sheet" or not adopted_batch(record)]
        if len(repeats.get(core, ())) > 1 and word_candidates:
            group["status"] = "pending"
            group["reasons"].append("同日同廠同方法同點位有不同 B，保留為重採；Word 無 B 無法唯一配對")
            continue
        exact_site = [record for record in base_candidates if _text(record["current"].get("site")) == sheet["site"]]
        blank_site = [record for record in base_candidates if not _text(record["current"].get("site"))]
        candidates = exact_site if exact_site else blank_site if blank_site else base_candidates
        if len(candidates) > 1:
            compatible = [record for record in candidates if not mismatches(sheet, record["current"])]
            if compatible:
                candidates = compatible
            else:
                group["status"] = "conflict"
                group["reasons"].append("Word 日期／方法／點位候選皆與 Sheet 非空欄位衝突")
                continue
        if len(candidates) > 1:
            group["status"] = "pending"
            group["reasons"].append(f"Word 有 {len(candidates)} 筆日期／方法／點位候選，非唯一")
            continue
        if len(candidates) == 1:
            record = candidates[0]
            current = record["current"]
            differences = mismatches(sheet, current)
            if differences:
                group["status"] = "conflict"
                group["reasons"].append("Word 非空欄位衝突：" + ", ".join(differences))
                continue
            fills = {field: value for field, value in sheet.items()
                     if value and not current.get(field) and not record.get("verified", {}).get(field)}
            group["word_record_id"] = record["id"]
            group["word_version"] = record["version"]
            group["suggested_fills"] = fills
            group["status"] = "supplement" if fills else "matched"
            continue
    return groups


def summarize(snapshot: dict, groups: list[dict]) -> dict:
    counts = {key: 0 for key in ("matched", "supplement", "conflict", "new", "pending", "already_adopted")}
    for group in groups:
        counts[group["status"]] = counts.get(group["status"], 0) + 1
    rows = [row for sheet_rows in snapshot["rows_by_sheet"].values() for row in sheet_rows
            if any(row["values"]) and not _is_header(row["values"])]
    return {"complete": bool(snapshot.get("complete")), "sheets": len(snapshot["rows_by_sheet"]),
            "source_rows": len(rows), "sample_groups": len(groups),
            "species_child_rows": sum(int(group["species_rows"]) for group in groups), **counts}


def missing_from_complete_snapshot(current_groups: list[dict], previous_snapshot: dict | None,
                                   *, current_complete: bool = False) -> list[dict]:
    """Report samples absent from the next complete snapshot without changing old records."""
    if not current_complete or not previous_snapshot or not previous_snapshot.get("complete"):
        return []
    current = {group["fingerprint"] for group in current_groups}
    missing = []
    for group in sample_groups(previous_snapshot):
        if (group["fingerprint"].startswith("pending:") or group["fingerprint"] in current or
                not group.get("sample")):
            continue
        item = dict(group)
        item["status"] = "missing"
        item["reasons"] = ["此完整快照未見較早完整快照中的採樣；只提示，不作廢或刪除紀錄"]
        missing.append(item)
    return missing


READONLY_SCOPE = "https://www.googleapis.com/auth/spreadsheets.readonly"
OAUTH_AUTHORIZE = "https://accounts.google.com/o/oauth2/v2/auth"
OAUTH_TOKEN = "https://oauth2.googleapis.com/token"
SHEETS_API = "https://sheets.googleapis.com/v4/spreadsheets/"


def google_oauth_status(data_dir):
    configured = bool(os.environ.get("EM_MVP_GOOGLE_CLIENT_ID") and
                      os.environ.get("EM_MVP_GOOGLE_CLIENT_SECRET"))
    token_path = Path(data_dir) / "google-sheet-token.json"
    authorized = False
    if token_path.is_file():
        try:
            token = json.loads(token_path.read_text(encoding="utf-8"))
            authorized = bool(token.get("refresh_token"))
        except (OSError, ValueError):
            authorized = False
    return {"configured": configured, "authorized": configured and authorized}


def google_authorize_url(redirect_uri, state):
    client_id = os.environ.get("EM_MVP_GOOGLE_CLIENT_ID", "").strip()
    if not client_id or not os.environ.get("EM_MVP_GOOGLE_CLIENT_SECRET"):
        raise ValueError("請先在本機環境設定 EM_MVP_GOOGLE_CLIENT_ID 與 EM_MVP_GOOGLE_CLIENT_SECRET。")
    query = urlencode({"client_id": client_id, "redirect_uri": redirect_uri,
                       "response_type": "code", "scope": READONLY_SCOPE,
                       "access_type": "offline", "prompt": "consent", "state": state})
    return OAUTH_AUTHORIZE + "?" + query


def _oauth_post(form):
    request = Request(OAUTH_TOKEN, data=urlencode(form).encode("ascii"),
                      headers={"Content-Type": "application/x-www-form-urlencoded"}, method="POST")
    try:
        with urlopen(request, timeout=20) as response:
            return json.loads(response.read().decode("utf-8"))
    except (HTTPError, URLError, TimeoutError, ValueError) as error:
        status = getattr(error, "code", "connection")
        raise ValueError(f"Google OAuth 未完成（{status}）；請檢查本機用戶端設定並重試。") from error


def _save_google_token(data_dir, token):
    path = Path(data_dir) / "google-sheet-token.json"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(token, ensure_ascii=False), encoding="utf-8")
    try:
        path.chmod(0o600)
    except OSError:
        pass  # Windows ACL inheritance governs access on the local profile.


def finish_google_oauth(data_dir, code, redirect_uri):
    client_id = os.environ.get("EM_MVP_GOOGLE_CLIENT_ID", "").strip()
    client_secret = os.environ.get("EM_MVP_GOOGLE_CLIENT_SECRET", "")
    if not client_id or not client_secret:
        raise ValueError("本機 Google OAuth 用戶端設定不完整。")
    response = _oauth_post({"code": code, "client_id": client_id, "client_secret": client_secret,
                            "redirect_uri": redirect_uri, "grant_type": "authorization_code"})
    refresh_token = response.get("refresh_token")
    if not refresh_token:
        raise ValueError("Google 未回傳離線唯讀權杖；未儲存授權。請重新授權並同意離線存取。")
    response["refresh_token"] = refresh_token
    response["client_id"] = client_id
    response["token_uri"] = OAUTH_TOKEN
    _save_google_token(data_dir, response)


def _google_access_token(data_dir):
    status = google_oauth_status(data_dir)
    if not status["configured"]:
        raise ValueError("Google Sheet 直連尚未設定本機 OAuth 用戶端；可先用兩分頁 XLSX 路徑。")
    path = Path(data_dir) / "google-sheet-token.json"
    try:
        token = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError) as error:
        raise ValueError("尚未完成 Google Sheet 唯讀授權；請先按「連接 Google Sheet」。") from error
    if not token.get("refresh_token"):
        raise ValueError("本機 OAuth 權杖缺少 refresh token；請重新授權。")
    expires_at = float(token.get("expires_at", 0))
    if token.get("access_token") and expires_at > datetime.now().timestamp() + 60:
        return token["access_token"]
    response = _oauth_post({"client_id": os.environ["EM_MVP_GOOGLE_CLIENT_ID"],
                            "client_secret": os.environ["EM_MVP_GOOGLE_CLIENT_SECRET"],
                            "refresh_token": token["refresh_token"], "grant_type": "refresh_token"})
    token.update(response)
    token["expires_at"] = datetime.now().timestamp() + int(response.get("expires_in", 3600))
    _save_google_token(data_dir, token)
    return token["access_token"]


def _google_get(url, token):
    request = Request(url, headers={"Authorization": "Bearer " + token, "Accept": "application/json"})
    try:
        with urlopen(request, timeout=30) as response:
            return json.loads(response.read().decode("utf-8"))
    except (HTTPError, URLError, TimeoutError, ValueError) as error:
        status = getattr(error, "code", "connection")
        raise ValueError(f"Google Sheets 唯讀請求失敗（{status}）；快照未寫入。") from error


def read_google_targets(data_dir):
    """Fetch metadata then bounded A:Y chunks from exactly the two configured tabs."""
    token = _google_access_token(data_dir)
    metadata = _google_get(SHEETS_API + SPREADSHEET_ID + "?" + urlencode({
        "fields": "sheets(properties(sheetId,title,gridProperties(rowCount,columnCount)))"}), token)
    sheet_props = {item["properties"]["title"]: item["properties"] for item in metadata.get("sheets", [])}
    for title, expected in SHEETS.items():
        props = sheet_props.get(title, {})
        if props.get("sheetId") != expected["sheet_id"] or int(props.get("gridProperties", {}).get("columnCount", 0)) < 25:
            raise ValueError(f"Google Sheet 的指定分頁 metadata 與工單不符：{title}；快照未寫入。")
    ranges = {}
    tab_evidence = {}
    for title in SHEETS:  # Deliberately never request another tab's values.
        total = int(sheet_props[title]["gridProperties"]["rowCount"])
        tab_evidence[title] = {"sheet_id": SHEETS[title]["sheet_id"], "range": f"A1:Y{total}",
                              "row_count": total}
        rows = []
        for start in range(1, total + 1, 400):
            end = min(total, start + 399)
            range_name = f"'{title}'!A{start}:Y{end}"
            url = SHEETS_API + SPREADSHEET_ID + "/values/" + __import__("urllib.parse").parse.quote(range_name, safe="")
            payload = _google_get(url + "?" + urlencode({"valueRenderOption": "UNFORMATTED_VALUE",
                                                           "majorDimension": "ROWS"}), token)
            chunk = payload.get("values", [])
            rows.extend(chunk)
            rows.extend([[]] * max(0, end - start + 1 - len(chunk)))
        ranges[title] = rows
    return parse_api_ranges(ranges, complete=True, evidence={
        "source_kind": "google-api", "spreadsheet_id": SPREADSHEET_ID,
        "completeness_basis": "Google Sheets API metadata rowCount 全範圍分段讀取",
        "tabs": tab_evidence,
    })
