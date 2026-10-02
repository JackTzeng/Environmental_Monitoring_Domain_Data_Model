"""Read source files without changing them. Uncertain values stay explicit."""
from __future__ import annotations

import csv
import io
import re
import unicodedata
import zipfile
from datetime import date, datetime
from pathlib import Path
from xml.etree import ElementTree as ET

from .domain import FIELDS, normalize_method
from .point_room_maps import POINT_ROOMS, ROOM_MAP_EVIDENCE, ROOM_MAP_SCOPE

PARSER_VERSION = "0.1.3"
MAX_DOCX_XML_BYTES = 20_000_000
MAX_XLSX_UNCOMPRESSED_BYTES = 80_000_000
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
W = "{" + NS["w"] + "}"


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def _key(value: object) -> str:
    return re.sub(r"[\s:：()（）_/.-]+", "", unicodedata.normalize("NFKC", _text(value))).lower()


_ALIASES = {
    "sample_date": ("監測日期", "採樣日期", "檢測日期", "日期", "sample_date", "date"),
    "site": ("廠別", "廠區", "site", "factory"),
    "monitoring_type": ("監測類型", "檢測目的", "類型", "monitoring_type", "purpose"),
    "method": ("監測方式", "檢測方式", "採樣方法", "採樣方式", "方法", "method"),
    "room": ("房間", "房間編號", "科室", "區域", "room"),
    "grade": ("Grade", "等級", "潔淨等級"),
    "point_id": ("監測位置", "點位", "採樣點位", "點位編號", "point_id", "point"),
    "operator": ("操作者", "操作人員", "操作人", "operator"),
    "batch_no": ("檢驗批次", "批號", "批次", "批次編碼", "產品批號", "batch_no", "batch"),
    "result_raw": ("結果", "結果CFU", "總菌數", "總菌數CFU", "CFU", "result_raw", "菌落數", "菌落數CFU"),
    "unit": ("單位", "unit"),
    "organism_name": ("菌種名稱", "菌種", "菌相", "organism_name"),
    "alert_raw": ("Alarm Lv", "Alarm Lv CFU", "Alert", "警戒限值", "警戒值", "alert_raw"),
    "action_raw": ("Action Lv", "Action Lv CFU", "Action", "行動限值", "行動值", "action_raw"),
    "_year": ("年", "year"), "_month": ("月", "month"), "_day": ("日", "day"),
    "_checked": ("核", "核閱", "核准", "checked"),
}
_HEADER_FIELDS = {_key(alias): field for field, aliases in _ALIASES.items() for alias in aliases}
_CHECKED_BOXES = set("▓■▣☑☒✓✔")
_BOX_TOKEN = re.compile(r"([▓■▣☑☒✓✔□☐▢])\s*([^▓■▣☑☒✓✔□☐▢]+)")


def _header_field(value: object) -> str:
    key = _key(value)
    if key in _HEADER_FIELDS:
        return _HEADER_FIELDS[key]
    # Labels with extra unit descriptions are supported; unrelated numeric columns are not.
    if key.startswith(("結果", "總菌數")):
        return "result_raw"
    return ""


def _iso_date(value: str) -> str:
    match = re.fullmatch(r"(\d{4})[-/年](\d{1,2})[-/月](\d{1,2})(?:日)?", value.strip())
    if not match:
        return ""
    try:
        return date(*(int(part) for part in match.groups())).isoformat()
    except ValueError:
        return ""


def _result(raw: str) -> tuple[str, str]:
    text = unicodedata.normalize("NFKC", raw).strip()
    if not text:
        return "missing", ""
    if re.fullmatch(r"\d+(?:\s*CFU(?:\s*/\s*(?:plate|皿))?)?", text, re.I):
        return "count", str(int(re.match(r"\d+", text).group()))
    if re.fullmatch(r"TNTC|too numerous to count|不可計數|無法計數", text, re.I):
        return "tntc", ""
    if re.fullmatch(r"<\s*\d+(?:\.\d+)?(?:\s*CFU)?", text, re.I):
        return "less_than", ""
    if re.fullmatch(r"N/?A|不適用", text, re.I):
        return "not_applicable", ""
    return "unknown", ""


def _selected_box(value: str, field: str, warnings: list[str]) -> str:
    tokens = _BOX_TOKEN.findall(value)
    selected = [label.strip() for mark, label in tokens if mark in _CHECKED_BOXES]
    if len(selected) == 1:
        return selected[0]
    warnings.append(f"{field}: 勾選數量為 {len(selected)}，保留待核對")
    return ""


def _metadata(filename: str, texts: list[str], context: dict | None = None) -> dict:
    metadata: dict[str, str] = {}
    site = re.search(r"(?:一廠|三廠|1廠|3廠)", filename)
    if site:
        metadata["site"] = {"1廠": "一廠", "3廠": "三廠"}.get(site.group(), site.group())
    explicit_site = metadata.get("site", "")
    exceptional_purpose = next((word for word in ("污染", "汙染", "重測", "超標", "無塵衣", "製程人員") if word in filename), "")
    if exceptional_purpose:
        metadata["monitoring_type"] = exceptional_purpose
    elif "製程監控" in filename or "製程監測" in filename:
        metadata["monitoring_type"] = "製程監測"
    elif "環境微生物監" in filename:
        metadata["monitoring_type"] = "例行環測"
    file_date = re.search(r"(?<!\d)(20\d{2})(\d{2})(\d{2})(?!\d)", filename)
    if file_date:
        metadata["sample_date"] = _iso_date("-".join(file_date.groups()))
    for text in texts:
        match = re.fullmatch(r"\s*([^:：]{1,40})[:：]\s*(.*?)\s*", text, re.S)
        if not match:
            continue
        field = _header_field(match.group(1))
        if field in {"sample_date", "site", "monitoring_type", "method", "room", "operator", "batch_no"}:
            value = match.group(2).strip()
            if field in {"method", "monitoring_type"} and _BOX_TOKEN.search(value):
                box_warnings: list[str] = []
                value = _selected_box(value, field, box_warnings)
                metadata.setdefault("_warnings", []).extend(box_warnings)
            if field == "site" and value:
                value = {"1廠": "一廠", "3廠": "三廠"}.get(value, value)
                explicit_site = value
            metadata[field] = value
    site_context = (context or {}).get("site", "") if isinstance(context, dict) else ""
    if explicit_site and site_context and explicit_site != site_context:
        metadata.setdefault("_warnings", []).append(
            f"site: 文件廠別 {explicit_site} 與匯入資料夾脈絡 {site_context} 不同；保留文件廠別")
    elif not explicit_site and site_context:
        metadata["site"] = site_context
        metadata.setdefault("_warnings", []).append(
            f"site: 文件未提供廠別；依匯入資料夾脈絡補為 {site_context}")
    return metadata


def _make_record(values: dict[str, str], location: str, source: str, warnings: list[str]) -> dict:
    warnings.extend(values.pop("_warnings", []))
    values = {field: _text(values.get(field, "")) for field in FIELDS}
    if values["sample_date"]:
        raw_date = values["sample_date"]
        values["sample_date"] = _iso_date(raw_date)
        if not values["sample_date"]:
            warnings.append(f"sample_date: 無法確認有效日期 {raw_date!r}")
    values["method"] = normalize_method(values["method"])
    position = values["point_id"]
    method_position = re.fullmatch(r"(落菌|落菌法|指壓|培養皿接觸法|表面接觸|空氣採樣|空氣取樣)\s*[:：]\s*(.+)", position)
    if method_position:
        if not values["method"]:
            values["method"] = method_position.group(1)
        values["point_id"] = method_position.group(2)
        point_operator = re.search(r"操作後\s*([A-Za-z]{2,5})(?=\s|[-（(]|$)", method_position.group(2))
        if point_operator and not values["operator"]:
            values["operator"] = point_operator.group(1).upper()
    values["method"] = normalize_method(values["method"])
    point_room = re.fullmatch(r"([A-Za-z]\d{1,3}[A-Za-z]?|\d{1,3})\s*[:：]\s*(C\d{2}[A-Za-z]?)(?![A-Za-z0-9])(.*)", position, re.I)
    room_point = re.fullmatch(r"(C\d{2}[A-Za-z]?)\s*[:：]\s*([A-Za-z]\d{1,3}[A-Za-z]?|\d{1,3})", position, re.I)
    labeled_room_point = re.fullmatch(r"(C\d{2}[A-Za-z]?)_[^:：]+[:：]\s*([A-Za-z0-9][A-Za-z0-9._-]*(?:[（(][^）)]*[）)])?)", position, re.I)
    bsc_room_point = re.fullmatch(r"(C\d{2}[A-Za-z]?)_(BSC\d{2}-\d+(?:-\d+)?)", position, re.I)
    if bsc_room_point:
        room, point = bsc_room_point.groups()
    elif point_room or room_point or labeled_room_point:
        point, room = point_room.groups()[:2] if point_room else tuple(reversed((room_point or labeled_room_point).groups()))
    else:
        point = room = ""
    if point or room:
        values["point_id"] = point
        if values["room"] and values["room"].upper() != room.upper():
            warnings.append("room: 點位原文與房間欄位不一致，保留房間欄位供核對")
        elif not values["room"]:
            values["room"] = room.upper()
    mapped = re.fullmatch(r"(GTP7|GTP8)\s*[：:]\s*(BSC\d{2}(?:-\d+){1,2})", position, re.I)
    if mapped:
        factory_prefix, point = mapped.group(1).upper(), mapped.group(2).upper()
        expected_point_prefix = {"GTP7": "BSC07", "GTP8": "BSC08"}[factory_prefix]
        if not point.startswith(expected_point_prefix):
            warnings.append(f"point_id: {factory_prefix} 與 {point[:5]} 點位前綴不一致，待核對")
        else:
            values["point_id"] = point
            sample_year = values["sample_date"][:4] if values["sample_date"] else ""
            in_scope = (values["site"] == ROOM_MAP_SCOPE["site"] and
                        values["method"] in ("落菌", "落菌法") and sample_year == str(ROOM_MAP_SCOPE["year"]))
            reference_room = POINT_ROOMS.get(point) if in_scope else None
            if reference_room and not values["room"]:
                values["room"] = reference_room
                warnings.append(f"room: {ROOM_MAP_EVIDENCE}；適用範圍為三廠／落菌法／2026")
            elif reference_room and values["room"].upper() != reference_room:
                warnings.append(f"room: 來源明確填寫 {values['room']}，與參考對照 {reference_room} 不同；保留來源值待核對")
            elif reference_room:
                warnings.append(f"room: {ROOM_MAP_EVIDENCE}；來源房間與對照相符")
            else:
                warnings.append(f"room: GTP點位 {point} 不符合已核對對照範圍或無明確對照，待對照")
    if any(word in _key(position) for word in ("negativecontrol", "positivecontrol", "陰性對照", "陽性對照")):
        warnings.append("point_id: 來源為對照組，需確認是否應納入監測趨勢")
    purpose = values["monitoring_type"]
    if purpose in {"環測", "例行環測", "例行環境監測"}:
        values["monitoring_type"] = "例行環測"
    elif purpose in {"製程監測", "製程監控", "製測"}:
        values["monitoring_type"] = "製程監測"
    elif purpose == "無塵衣月監測":
        values["monitoring_type"] = "無塵衣月監測"
    elif purpose == "無塵衣季監測":
        values["monitoring_type"] = "無塵衣季監測"
    elif "無塵衣" in purpose:
        values["monitoring_type"] = purpose
        warnings.append(f"monitoring_type: 保留來源類型 {purpose!r}，需確認")
    else:
        values["monitoring_type"] = "待分類"
        warnings.append(f"monitoring_type: 需確認原始目的 {purpose!r}")
    kind, count = _result(values["result_raw"])
    values["result_type"], values["cfu_count"] = kind, count
    if not values["unit"] and re.search(r"\bCFU\b", values["result_raw"], re.I):
        values["unit"] = "CFU"
    if any(item.startswith("result_raw: 多個") for item in warnings):
        values["result_type"], values["cfu_count"] = "unknown", ""
    if kind != "count":
        warnings.append(f"result_raw: {kind}，不轉成一般菌落數")
    for field in FIELDS:
        if not values[field] and field not in {"cfu_count"}:
            warnings.append(f"{field}: 未取得")
    return {"values": values, "source_location": location, "source_text": source,
            "warnings": list(dict.fromkeys(warnings))}


def _find_header(rows: list[list[str]]) -> tuple[int, dict[str, list[int]]]:
    for index, row in enumerate(rows[:40]):
        mapping: dict[str, list[int]] = {}
        for column, label in enumerate(row):
            field = _header_field(label)
            if field:
                mapping.setdefault(field, []).append(column)
        if "result_raw" in mapping and any(field in mapping for field in ("point_id", "operator", "room", "method")):
            return index, mapping
    raise ValueError("找不到可辨識的點位/操作者及結果欄位；請先核對來源版型。")


def _tabular_records(rows: list[list[str]], name: str, metadata: dict[str, str], *, sheet: bool) -> list[dict]:
    header_index, mapping = _find_header(rows)
    metadata = dict(metadata)
    headings = " ".join(cell for row in rows[:header_index + 1] for cell in row)
    sites = set(re.findall(r"一廠|三廠|1廠|3廠", headings))
    if not metadata.get("site") and len(sites) == 1:
        site = sites.pop()
        metadata["site"] = {"1廠": "一廠", "3廠": "三廠"}.get(site, site)
    records = []
    header = rows[header_index]
    for row_index, row in enumerate(rows[header_index + 1:], start=header_index + 2):
        if not any(cell.strip() for cell in row):
            continue
        if row == header or any(row[0].startswith(label) for label in ("製表", "核閱", "以下空白", "以下空格")):
            continue
        warnings: list[str] = list(metadata.get("_warnings", []))
        values = dict(metadata)
        for field, columns in mapping.items():
            if field.startswith("_"):
                continue
            actual = [(column, row[column] if column < len(row) else "") for column in columns]
            distinct = list(dict.fromkeys(value for _, value in actual if value.strip()))
            if field == "result_raw" and len(columns) > 1:
                if sheet and any(_key(header[column]).startswith("總菌數") for column in columns):
                    total_column = next(column for column in columns if _key(header[column]).startswith("總菌數"))
                    values[field] = row[total_column] if total_column < len(row) else ""
                    if len(distinct) > 1:
                        warnings.append("result_raw: 總菌數與菌落數不同；本列採總菌數，菌種列可能重复，需核對")
                else:
                    values[field] = " | ".join(value for _, value in actual)
                    warnings.append("result_raw: 多個結果欄位，未猜測應用哪個結果")
            elif len(distinct) > 1:
                values[field] = ""
                warnings.append(f"{field}: 多個欄位值互相衝突")
            else:
                values[field] = distinct[0] if distinct else values.get(field, "")
        if all(field in mapping for field in ("_year", "_month", "_day")):
            parts = [row[mapping[field][0]] if mapping[field][0] < len(row) else "" for field in ("_year", "_month", "_day")]
            values["sample_date"] = "-".join(parts)
        if "_checked" in mapping:
            column = mapping["_checked"][0]
            checked = row[column] if column < len(row) else ""
            if checked.upper() != "TRUE":
                warnings.append(f"來源核閱欄不是 TRUE ({checked!r})；需人工確認是否納入")
        if not values.get("unit") and any("cfu" in _key(header[column]) for column in mapping["result_raw"]):
            values["unit"] = "CFU"
        source_text = " | ".join(row)
        point = values.get("point_id", "").strip()
        sample_keys = ("point_id", "room", "operator", "batch_no", "result_raw")
        has_sample_data = any(
            column < len(row) and row[column].strip()
            for field in sample_keys for column in mapping.get(field, [])
        )
        footer = any(re.search(r"以下空白|以下空格", cell) for cell in row)
        judgement_only = any(re.search(r"□\s*合格\s*□\s*不合格", cell) for cell in row)
        if not has_sample_data and (footer or judgement_only):
            continue
        if point and _key(point) in {"以下空白", "以下空格"} and not values.get("result_raw", "").strip():
            continue
        for column, label in enumerate(header):
            if _key(label).startswith("標準") and column < len(row) and row[column].strip():
                warnings.append(f"standard: 來源標準(CFU)={row[column].strip()}；保留原文，不轉為 Alert/Action")
        records.append(_make_record(values, f"{name} row {row_index}", " | ".join(row), warnings))
    return records


def _word(filename: str, data: bytes, context: dict | None = None) -> list[dict]:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            info = archive.getinfo("word/document.xml")
            if info.file_size > MAX_DOCX_XML_BYTES:
                raise ValueError("Word 文件 XML 過大，無法匯入。")
            xml = archive.read(info)
        if b"<!DOCTYPE" in xml or b"<!ENTITY" in xml:
            raise ValueError("Word 含不支援的 XML 宣告。")
        document = ET.fromstring(xml)
    except (zipfile.BadZipFile, KeyError, ET.ParseError) as exc:
        raise ValueError("無法讀取有效的 .docx 文件。") from exc
    texts = ["".join(node.itertext()).strip() for node in document.findall(".//w:body/w:p", NS)]
    # Metadata in key/value control tables is read only when the following cell is a value.
    for table in document.findall(".//w:body/w:tbl", NS):
        for row in table.findall("./w:tr", NS):
            cells = [_cell_text(cell) for cell in row.findall("./w:tc", NS)]
            texts.extend(cells)
            for label, value in zip(cells, cells[1:]):
                if _header_field(label) in {"sample_date", "site", "method", "room", "operator", "batch_no", "monitoring_type"} and not _header_field(value):
                    texts.append(label + "：" + value)
    metadata = _metadata(filename, texts, context)
    records: list[dict] = []
    for table_index, table in enumerate(document.findall(".//w:body/w:tbl", NS), start=1):
        rows: list[list[str]] = []
        previous: dict[int, str] = {}
        for row in table.findall("./w:tr", NS):
            values: list[str] = []
            next_previous: dict[int, str] = {}
            for cell in row.findall("./w:tc", NS):
                text = _cell_text(cell)
                span_node = cell.find("./w:tcPr/w:gridSpan", NS)
                span = int(span_node.get(W + "val", "1")) if span_node is not None else 1
                if not 1 <= span <= 100:
                    raise ValueError("Word 合併欄位範圍不合法。")
                merge = cell.find("./w:tcPr/w:vMerge", NS)
                if merge is not None and merge.get(W + "val") != "restart" and not text:
                    text = previous.get(len(values), "")
                for _ in range(span):
                    if merge is not None:
                        next_previous[len(values)] = text
                    values.append(text)
            rows.append(values)
            previous = next_previous
        try:
            records.extend(_tabular_records(rows, f"Word table {table_index}", metadata, sheet=False))
        except ValueError as exc:
            if not str(exc).startswith("找不到可辨識"):
                raise
    if not records:
        raise ValueError("Word 沒有可辨識的監測明細；請核對表格標頭及資料列。")
    return records


def _cell_text(cell: ET.Element) -> str:
    return "\n".join("".join(node.text or "" for node in paragraph.findall(".//w:t", NS))
                     for paragraph in cell.findall("./w:p", NS)).strip()


def parse_file(filename: str, data: bytes, context: dict | None = None) -> list[dict]:
    """Return rows with raw evidence. No source file or database is written."""
    extension = Path(filename).suffix.lower()
    if not data:
        raise ValueError("來源檔案是空的。")
    if extension == ".docx":
        return _word(filename, data, context)
    metadata = _metadata(filename, [], context)
    if extension == ".csv":
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError("CSV 必須使用 UTF-8 編碼；請從 Google Sheet/Excel 重新匯出。") from exc
        rows = [[_text(value) for value in row] for row in csv.reader(io.StringIO(text, newline=""))]
        return _tabular_records(rows, "CSV", metadata, sheet=True)
    if extension == ".xlsx":
        from openpyxl import load_workbook
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(item.file_size for item in archive.infolist()) > MAX_XLSX_UNCOMPRESSED_BYTES:
                    raise ValueError("Excel 解壓縮後資料超過 80 MB，無法匯入。")
                if any(item.filename.lower().endswith("vbaproject.bin") for item in archive.infolist()):
                    raise ValueError("不支援含 VBA 巨集的 Excel 檔案；請另存為無巨集 .xlsx。")
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
        except (zipfile.BadZipFile, KeyError, ET.ParseError) as exc:
            raise ValueError("無法讀取有效的 .xlsx 檔案。") from exc
        records = []
        try:
            for worksheet in workbook.worksheets:
                rows = [[_text(value) for value in row] for row in worksheet.iter_rows(values_only=True)]
                try:
                    parsed = _tabular_records(rows, "Sheet " + worksheet.title, metadata, sheet=True)
                    for record in parsed:
                        if any(value.startswith("=") for value in record["values"].values()):
                            record["warnings"].append("來源含 Excel 公式；請先核對公式結果後輸入確定值")
                    records.extend(parsed)
                except ValueError as exc:
                    if not str(exc).startswith("找不到可辨識"):
                        raise
        finally:
            workbook.close()
        if not records:
            raise ValueError("Excel 沒有可辨識的監測明細。")
        return records
    raise ValueError("只支援 .docx、.xlsx、UTF-8 .csv 檔案。")
